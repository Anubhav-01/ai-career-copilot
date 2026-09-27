"""File validation and text extraction for resume uploads (PDF / DOCX).

Validation is defense-in-depth: extension allowlist, size limit, and magic
byte checks before any parser touches the file.
"""
import io

from app.core.errors import FileUploadError, TextExtractionError

ALLOWED_EXTENSIONS = {"pdf", "docx"}

_PDF_MAGIC = b"%PDF"
_ZIP_MAGIC = b"PK\x03\x04"  # docx is a zip container


def validate_upload(filename: str, content: bytes, max_size_bytes: int) -> str:
    """Validate an uploaded file; returns the normalized extension."""
    if not filename or "." not in filename:
        raise FileUploadError("File must have a .pdf or .docx extension.")
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise FileUploadError(f"Unsupported file type '.{ext}'. Upload a PDF or DOCX.")
    if len(content) == 0:
        raise FileUploadError("The uploaded file is empty.")
    if len(content) > max_size_bytes:
        raise FileUploadError(
            f"File is too large ({len(content) // 1024} KB)."
            f" Maximum allowed is {max_size_bytes // (1024 * 1024)} MB."
        )
    if ext == "pdf" and not content.startswith(_PDF_MAGIC):
        raise FileUploadError("This file does not look like a valid PDF.")
    if ext == "docx" and not content.startswith(_ZIP_MAGIC):
        raise FileUploadError("This file does not look like a valid DOCX.")
    return ext


def extract_text(content: bytes, ext: str) -> str:
    if ext == "pdf":
        text = _extract_pdf(content)
    elif ext == "docx":
        text = _extract_docx(content)
    else:
        raise FileUploadError(f"Unsupported file type '.{ext}'.")

    cleaned = "\n".join(line.rstrip() for line in text.splitlines())
    if len(cleaned.strip()) < 50:
        raise TextExtractionError(
            "Could not extract readable text. If this is a scanned/image PDF,"
            " export a text-based version and try again."
        )
    return cleaned.strip()


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise TextExtractionError("Encrypted PDFs are not supported.")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except TextExtractionError:
        raise
    except Exception as exc:
        raise TextExtractionError("Failed to read this PDF file.") from exc


def _extract_docx(content: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(content))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise TextExtractionError("Failed to read this DOCX file.") from exc
